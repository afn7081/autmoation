"""
Generates personalized certificate PNGs from a PPTX template.

Flow:
  1. Open the PPTX template
  2. Replace "Full Name" and date placeholder with donor-specific values
  3. Convert PPTX -> PNG. By default uses LibreOffice locally (no API limits).
     Set USE_LOCAL_PPTX_CONVERSION=false to fall back to CloudConvert.
  4. Return the PNG bytes

Env vars:
  - USE_LOCAL_PPTX_CONVERSION: "true" (default) or "false"
  - CLOUDCONVERT_API_KEY: required only when USE_LOCAL_PPTX_CONVERSION=false
"""

import os
import uuid
import base64
import shutil
import time
import logging
import platform
import tempfile
import subprocess
from datetime import datetime, timezone

import requests
from pptx import Presentation

from asset_store import get_asset

CERTIFICATE_TEMPLATE_BLOB = "cert_template_new.pptx"


def _use_local() -> bool:
    return os.getenv("USE_LOCAL_PPTX_CONVERSION", "true").strip().lower() in (
        "1", "true", "yes", "on",
    )


def _soffice_cmd() -> str:
    if platform.system() == "Windows":
        for p in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if os.path.exists(p):
                return p
        return "soffice"
    return "libreoffice"


def _replace_text_preserving_format(shape, old_text, new_text):
    for para in shape.text_frame.paragraphs:
        if old_text not in para.text:
            continue
        for run in para.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)
                return True
        if para.runs:
            combined = "".join(r.text for r in para.runs)
            if old_text in combined:
                para.runs[0].text = combined.replace(old_text, new_text)
                for r in para.runs[1:]:
                    r.text = ""
                return True
    return False


def _format_date(date: str) -> str:
    if not date:
        return datetime.now(timezone.utc).strftime("%m-%d-%y")
    try:
        return datetime.strptime(date, "%B %d, %Y").strftime("%m-%d-%y")
    except ValueError:
        return date


def _convert_pptx_to_png_local(pptx_path: str, work_dir: str) -> bytes:
    """Use LibreOffice headless to convert a PPTX to PNG (first slide)."""
    soffice = _soffice_cmd()
    user_profile = os.path.join(work_dir, f"libreoffice_profile_{uuid.uuid4().hex}")

    env = os.environ.copy()
    env["HOME"] = "/tmp"

    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--nofirststartwizard",
        f"-env:UserInstallation=file://{user_profile}",
        "--convert-to", "png",
        "--outdir", work_dir,
        pptx_path,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=180, env=env)

    base = os.path.splitext(os.path.basename(pptx_path))[0]
    png_path = os.path.join(work_dir, f"{base}.png")
    if not os.path.exists(png_path):
        raise RuntimeError(
            f"LibreOffice PPTX->PNG failed (exit {result.returncode}):\n"
            f"stdout: {result.stdout}\nstderr: {result.stderr}\ncmd: {' '.join(cmd)}"
        )
    with open(png_path, "rb") as f:
        return f.read()


def _convert_pptx_to_png_cloudconvert(pptx_path: str) -> bytes:
    """Convert a PPTX file to PNG bytes via CloudConvert.

    Requires CLOUDCONVERT_API_KEY (a Bearer token) in the environment.
    Creates a job with import-upload -> convert -> export-url tasks, uploads the
    PPTX, polls until the job finishes, and downloads the resulting PNG.
    """
    api_key = os.getenv("CLOUDCONVERT_API_KEY")
    if not api_key:
        raise ValueError("CLOUDCONVERT_API_KEY is required when USE_LOCAL_PPTX_CONVERSION=false")

    base_url = "https://api.cloudconvert.com/v2"
    headers = {"Authorization": f"Bearer {api_key}"}

    job_payload = {
        "tasks": {
            "import-pptx": {"operation": "import/upload"},
            "convert-pptx": {
                "operation": "convert",
                "input": "import-pptx",
                "input_format": "pptx",
                "output_format": "png",
            },
            "export-png": {"operation": "export/url", "input": "convert-pptx"},
        }
    }

    resp = requests.post(f"{base_url}/jobs", json=job_payload, headers=headers, timeout=60)
    if resp.status_code not in (200, 201):
        raise RuntimeError(f"CloudConvert job create failed ({resp.status_code}): {resp.text}")
    job = resp.json()["data"]
    job_id = job["id"]

    upload_task = next(t for t in job["tasks"] if t["name"] == "import-pptx")
    form = upload_task["result"]["form"]
    with open(pptx_path, "rb") as f:
        up = requests.post(
            form["url"],
            data=form["parameters"],
            files={"file": (
                "certificate.pptx",
                f,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )},
            timeout=120,
        )
    if up.status_code not in (200, 201, 204):
        raise RuntimeError(f"CloudConvert upload failed ({up.status_code}): {up.text}")

    deadline = time.time() + 180
    while time.time() < deadline:
        s = requests.get(f"{base_url}/jobs/{job_id}", headers=headers, timeout=30)
        s.raise_for_status()
        data = s.json()["data"]
        status = data["status"]
        if status == "finished":
            export_task = next(t for t in data["tasks"] if t["name"] == "export-png")
            files = (export_task.get("result") or {}).get("files") or []
            if not files:
                raise RuntimeError(f"CloudConvert export returned no files: {export_task}")
            r = requests.get(files[0]["url"], timeout=120)
            r.raise_for_status()
            return r.content
        if status == "error":
            raise RuntimeError(f"CloudConvert job error: {data}")
        time.sleep(2)

    raise RuntimeError("CloudConvert job timed out after 180s")


def generate_certificate_png(donor_name: str, date: str = "") -> bytes:
    template_path = get_asset(CERTIFICATE_TEMPLATE_BLOB)

    file_size = os.path.getsize(template_path)
    logging.info(f"Template file: {template_path} ({file_size} bytes)")

    formatted_date = _format_date(date)

    prs = Presentation(template_path)
    name_replaced = date_replaced = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if not name_replaced and _replace_text_preserving_format(shape, "Full Name", donor_name):
                name_replaced = True
            if not date_replaced and _replace_text_preserving_format(
                shape, "Date :  04-03-26", f"Date :  {formatted_date}"
            ):
                date_replaced = True

    if not name_replaced:
        logging.warning("PPTX: 'Full Name' placeholder not found")
    if not date_replaced:
        logging.warning("PPTX: date placeholder 'Date :  04-03-26' not found")

    work_dir = tempfile.mkdtemp(prefix="cert_")
    pptx_out = os.path.join(work_dir, "certificate.pptx")
    prs.save(pptx_out)
    logging.info(f"Certificate PPTX built for: {donor_name}")

    try:
        if _use_local():
            logging.info("Converting PPTX->PNG locally via LibreOffice")
            png_bytes = _convert_pptx_to_png_local(pptx_out, work_dir)
        else:
            logging.info("Converting PPTX->PNG via CloudConvert")
            png_bytes = _convert_pptx_to_png_cloudconvert(pptx_out)

        logging.info(f"Certificate PNG generated ({len(png_bytes)} bytes)")
        return png_bytes
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
